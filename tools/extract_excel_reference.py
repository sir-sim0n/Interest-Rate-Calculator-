"""Level 1 reference: the values Microsoft Excel saved in the workbook.

For every module this reads the current inputs and the cached outputs
(what Excel displayed when the file was last saved) and writes them to
``validation/fixtures/excel_saved_values.json``.

Run from the project root:   python tools/extract_excel_reference.py
"""

from __future__ import annotations

import json
import sys
import warnings
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from validation.excel_spec import MODULES, WORKBOOK_NAME  # noqa: E402

WORKBOOK = ROOT / "reference" / WORKBOOK_NAME
OUTPUT = ROOT / "validation" / "fixtures" / "excel_saved_values.json"


def clean(value):
    """JSON-friendly cell value: numbers as float, text/errors as str, blanks as None."""
    if value is None or isinstance(value, str):
        return value
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return float(value)
    return str(value)


def main() -> int:
    warnings.filterwarnings("ignore")
    wb = openpyxl.load_workbook(WORKBOOK, data_only=True)

    def get(sheet: str, address: str):
        return clean(wb[sheet][address].value)

    scenarios = []
    for module in MODULES:
        scenarios.append({
            "id": f"{module.name}/as_shipped",
            "module": module.name,
            "params": module.params_from_cells(get),
            "cells": {name: f"{module.cell(ref)[0]}!{module.cell(ref)[1]}" for name, ref in module.outputs.items()},
            "outputs": {name: get(*module.cell(ref)) for name, ref in module.outputs.items()},
        })

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT.write_text(json.dumps({
        "source": WORKBOOK_NAME,
        "engine": "Microsoft Excel (values saved in the workbook)",
        "scenarios": scenarios,
    }, indent=2, ensure_ascii=False), encoding="utf-8")
    print(f"wrote {len(scenarios)} scenarios -> {OUTPUT.relative_to(ROOT)}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
