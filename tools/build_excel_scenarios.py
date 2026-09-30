"""Level 2 reference: new input scenarios recalculated from the workbook itself.

Excel is not needed. For each scenario in ``validation/excel_spec.py`` this

1. writes the scenario's inputs into a copy of the original workbook
   (formulas untouched, so the workbook's own logic does the calculating);
2. recalculates all copies headlessly with LibreOffice Calc;
3. reads the recalculated outputs and stores them in
   ``validation/fixtures/excel_scenarios.json``.

LibreOffice reproduced every deterministic cell of the workbook exactly
(529 of 529 numeric cells that do not use RAND() or the live FX data type),
so it is a faithful stand-in for Excel here.

Requirements: LibreOffice (``soffice`` on PATH) and openpyxl.
Run from the project root:   python tools/build_excel_scenarios.py
"""

from __future__ import annotations

import json
import shutil
import subprocess
import sys
import tempfile
import warnings
from collections import defaultdict
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from validation.excel_spec import MODULES_BY_NAME, WORKBOOK_NAME, scenario_list  # noqa: E402
from tools.extract_excel_reference import clean  # noqa: E402

WORKBOOK = ROOT / "reference" / WORKBOOK_NAME
OUTPUT = ROOT / "validation" / "fixtures" / "excel_scenarios.json"
CHUNK = 40  # files per LibreOffice call


def soffice_binary() -> str:
    for name in ("soffice", "libreoffice"):
        path = shutil.which(name)
        if path:
            return path
    raise SystemExit("LibreOffice (soffice) was not found on PATH.")


def main() -> int:
    warnings.filterwarnings("ignore")
    soffice = soffice_binary()
    version = subprocess.run([soffice, "--version"], capture_output=True, text=True).stdout.strip()

    wb = openpyxl.load_workbook(WORKBOOK)  # formulas, not values
    for ws in wb.worksheets:  # pictures play no part in the calculations and
        ws._images = []       # openpyxl cannot save them more than once
    prices = {}
    teb = wb["BACKROOM2_TEB"]
    for col in range(1, teb.max_column + 1):
        if teb.cell(18, col).value:
            prices[teb.cell(18, col).value] = teb.cell(19, col).value

    groups: dict[str, list[dict]] = defaultdict(list)
    for scenario in scenario_list():
        if scenario["module"] == "broker":  # Excel copies the price from the dropdown
            params = scenario["params"]
            params["exchange_price"] = prices[params["exchange_company"]]
            params["country_price"] = prices[params["country_company"]]
        groups[scenario["file_group"]].append(scenario)

    work = Path(tempfile.mkdtemp(prefix="irc_scenarios_"))
    (work / "in").mkdir()
    (work / "out").mkdir()
    files = {}
    for index, (group, scenarios) in enumerate(groups.items()):
        writes = {}
        for scenario in scenarios:
            writes.update(MODULES_BY_NAME[scenario["module"]].to_cells(scenario["params"]))
        originals = {key: wb[key[0]][key[1]].value for key in writes}
        for (sheet, address), value in writes.items():
            wb[sheet][address].value = value
        path = work / "in" / f"s{index:04d}.xlsx"
        wb.save(path)
        for (sheet, address), value in originals.items():
            wb[sheet][address].value = value
        files[group] = path.name
    print(f"wrote {len(files)} workbook copies, recalculating with {version} ...")

    names = sorted(files.values())
    for start in range(0, len(names), CHUNK):
        batch = [str(work / "in" / n) for n in names[start:start + CHUNK]]
        subprocess.run([soffice, "--headless", "--calc", "--convert-to", "xlsx", "--outdir",
                        str(work / "out"), *batch], check=True, capture_output=True)
        print(f"  recalculated {min(start + CHUNK, len(names))}/{len(names)}")

    results = []
    for group, scenarios in groups.items():
        out = openpyxl.load_workbook(work / "out" / files[group], data_only=True, read_only=True)
        for scenario in scenarios:
            module = MODULES_BY_NAME[scenario["module"]]
            cells = {name: module.cell(ref) for name, ref in module.outputs.items()}
            results.append({
                "id": scenario["id"],
                "module": scenario["module"],
                "params": scenario["params"],
                "cells": {name: f"{s}!{a}" for name, (s, a) in cells.items()},
                "outputs": {name: clean(out[s][a].value) for name, (s, a) in cells.items()},
            })
        out.close()

    OUTPUT.write_text(json.dumps({
        "source": WORKBOOK_NAME,
        "engine": f"{version} (headless recalculation of the original workbook formulas)",
        "scenarios": results,
    }, indent=1, ensure_ascii=False), encoding="utf-8")
    shutil.rmtree(work, ignore_errors=True)
    print(f"wrote {len(results)} scenarios -> {OUTPUT.relative_to(ROOT)}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
